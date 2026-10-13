"""Workbook audit and causal transformations, retaining every original column."""
from dataclasses import dataclass
from pathlib import Path
import re

import numpy as np
import pandas as pd
from openpyxl.utils import get_column_letter

TARGET_PREFIX = 'Growth rates of Gross domestic product'
INFORMATION_CLASS = 'PSEUDO_REAL_TIME_NOT_VINTAGE_VERIFIED'


@dataclass
class Inputs:
    monthly_raw: pd.DataFrame
    quarterly_raw: pd.DataFrame
    monthly: pd.DataFrame
    quarterly: pd.DataFrame
    gdp: pd.Series
    inventory: pd.DataFrame
    transformations: pd.DataFrame
    audit: dict


def decumulate(series: pd.Series, frequency: str) -> pd.Series:
    """Calendar reindexing makes a missing predecessor a missing increment."""
    if series.index.has_duplicates:
        raise ValueError('Duplicate reference periods')
    periods = series.index if isinstance(series.index, pd.PeriodIndex) else series.index.to_period(frequency)
    if periods.has_duplicates:
        raise ValueError('Duplicate reference periods')
    s = pd.Series(series.to_numpy(), index=periods).reindex(pd.period_range(periods.min(), periods.max(), freq=frequency))
    result = s.groupby(s.index.year).diff()
    first = s.index.quarter == 1 if frequency == 'Q' else s.index.month == 1
    result.loc[first] = s.loc[first]
    return result.reindex(periods).set_axis(series.index)


def log_growth(s: pd.Series, lag: int) -> pd.Series:
    prior = s.shift(lag)
    return 100 * np.log(s.where(s.gt(0)) / prior.where(prior.gt(0)))


def price_index(s: pd.Series, basis: str) -> pd.Series:
    """Indices are relative prices, never absolute price levels."""
    log = 100 * np.log(s.where(s.gt(0)) / 100)
    if basis == 'MOM':
        return log
    if basis == 'YTD':
        return decumulate(log, 'M')
    raise ValueError('Unspecified index basis: ' + basis)


def gdp_growth(s: pd.Series) -> pd.Series:
    if not s.dropna().between(70, 140).all():
        raise ValueError('GDP index convention incompatible with workbook')
    return s - 100


def specification(label: str, frequency: str) -> tuple[str, str, int, str]:
    """Explicit research assumptions; never writes the official registry."""
    if frequency == 'Q':
        if label.startswith('Remittance'):
            return 'external_income', 'LOG_YOY', 90, 'Quarter-specific secondary income credits; not YTD: Q4 often below Q3. Not an exact remittance measure.'
        if label.startswith('Average Nominal Wages'):
            return 'wages', 'LOG_YOY', 30, 'Published wage average; may be YTD average; never de-cumulate an average. Nominal, not real wages.'
        if label.startswith('Investments in fixed assets'):
            return 'investment', 'INDEX_MINUS_100', 45, 'Published real growth index, likely YTD; retain growth concept without de-cumulating a ratio.'
        if label.startswith('Investment financed'):
            return 'investment', 'DECUM_LOG_YOY', 45, 'YTD nominal FDI-funded investment, inferred from annual reset and within-year accumulation; definition unverified.'
        if label.startswith('Total income'):
            return 'household_income', 'DECUM_LOG_YOY', 45, 'YTD nominal population income, annual reset and within-year accumulation; definition unverified.'
    if label.startswith('Industrial production'):
        return 'industry', 'DECUM_LOG_YOY', 30, 'Nominal YTD industrial output; annual resets and within-year accumulation. Not a real production index.'
    if label.startswith(('Construction', 'Retail', 'Wholesale')):
        return 'domestic_demand', 'DECUM_LOG_YOY', 30, 'Nominal YTD flow, annual reset and within-year accumulation.'
    if label.startswith('CPI'):
        return 'consumer_prices', 'INDEX_MOM_LOG', 5, 'Previous-month=100, despite misleading headline label Cumulative; matches official 2021 monthly CPI values.'
    if label.startswith('PPI'):
        return 'producer_prices', 'INDEX_MOM_LOG', 20, 'Previous-month=100; SIAT dataset 1581 and composite 2019 values corroborate basis. Nonpositive indices invalid.'
    if label.startswith(('Export', 'Import')):
        return 'external_trade', 'DECUM_LOG_YOY', 30, 'YTD nominal USD trade flow; annual reset and within-year accumulation.'
    if label.startswith('Gold ('):
        return 'financial', 'LOG_MOM', 0, 'USD gold price: monthly log return; workbook aggregation/source not supplied.'
    if label.startswith('M2'):
        return 'financial', 'LOG_YOY', 25, 'Broad money nominal stock; log YoY change.'
    if label.startswith('International reserves'):
        return 'financial', 'LOG_MOM', 15, 'USD reserve stock; monthly log change, gold valuation included where applicable; swings retained.'
    if label.startswith('Rendered market services'):
        return 'services', 'INDEX_MINUS_100', 30, 'Published previous-year index; retain supplied monthly/YTD growth concept; cannot de-cumulate a growth ratio.'
    if label.startswith('Small business'):
        return 'enterprises', 'DECUM_LOG_YOY', 30, 'Label ambiguous: counts reset in January and accumulate; interpreted as newly created entities YTD, not stock. Provisional.'
    if label.startswith('Operating enterprises'):
        return 'enterprises', 'LOG_YOY', 30, 'Operating entity stock; no January reset; structural level changes retained.'
    if label.startswith('Interest rates'):
        return 'financial', 'LEVEL', 30, 'Annual percentage rate level, training-standardized; no log of a rate.'
    raise ValueError('Unmapped predictor: ' + label)


def transform(s: pd.Series, rule: str, frequency: str):
    lag = 12 if frequency == 'M' else 4
    intermediate = decumulate(s, frequency) if rule == 'DECUM_LOG_YOY' else s.copy()
    if rule in ('DECUM_LOG_YOY', 'LOG_YOY'):
        clean = log_growth(intermediate, lag)
    elif rule == 'LOG_MOM':
        clean = log_growth(s, 1)
    elif rule == 'INDEX_MOM_LOG':
        clean = price_index(s, 'MOM')
    elif rule == 'INDEX_MINUS_100':
        clean = s - 100
    elif rule == 'LEVEL':
        clean = s.copy()
    else:
        raise ValueError(rule)
    return intermediate, clean.replace([np.inf, -np.inf], np.nan)


def _sheet(path: Path, sheet: str, frequency: str) -> pd.DataFrame:
    header = pd.read_excel(path, sheet_name=sheet, header=None, nrows=1).iloc[0].tolist()
    if pd.Index(header).has_duplicates or any(pd.isna(x) for x in header):
        raise ValueError('Missing/duplicate original headers: ' + sheet)
    f = pd.read_excel(path, sheet_name=sheet)
    dates = pd.PeriodIndex(pd.to_datetime(f.iloc[:, 0]), freq='M') if frequency == 'M' else pd.PeriodIndex(f.iloc[:, 0].astype(str), freq='Q')
    if dates.has_duplicates or not dates.is_monotonic_increasing:
        raise ValueError('Duplicate/unsorted dates: ' + sheet)
    f = f.iloc[:, 1:].apply(pd.to_numeric, errors='raise')
    f.index = dates
    if np.isinf(f.to_numpy()).any():
        raise ValueError('Infinite workbook cells')
    return f.reindex(pd.period_range(dates.min(), dates.max(), freq=frequency))


def read_workbook(path: Path, master: Path | None = None) -> Inputs:
    if pd.ExcelFile(path).sheet_names != ['Monthly_Data', 'Quarterly_Data']:
        raise ValueError('Unexpected workbook sheets')
    m = _sheet(path, 'Monthly_Data', 'M')
    q = _sheet(path, 'Quarterly_Data', 'Q')
    targets = [c for c in q if c.startswith(TARGET_PREFIX)]
    if len(targets) != 1:
        raise ValueError('Exactly one GDP target required')
    target = targets[0]
    gdp = gdp_growth(q[target])
    gdp.name = target
    predictors = q.drop(columns=target)
    if m.shape[1] != 91 or predictors.shape[1] != 5:
        raise ValueError(f'Completeness contract failed: {m.shape[1]} M, {predictors.shape[1]} Q')
    if pd.Index([*m, *predictors]).has_duplicates:
        raise ValueError('Predictor label duplicated across sheets')
    transformed = []; inventory = []; changes = []
    for sheet, frequency, frame in [('Monthly_Data', 'M', m), ('Quarterly_Data', 'Q', predictors)]:
        clean_frame = pd.DataFrame(index=frame.index)
        for position, label in enumerate(frame, start=2):
            raw = frame[label]
            group, rule, lag, definition = specification(label, frequency)
            intermediate, clean = transform(raw, rule, frequency)
            clean_frame[label] = clean
            valid = raw.dropna(); usable = clean.dropna()
            units = re.findall(r'\(([^()]*)\)', label)
            unit = units[-1] if units else 'not supplied'
            annual = raw.groupby(raw.index.year).mean().pct_change(fill_method=None).abs()
            med = clean.median(); mad = (clean-med).abs().median()
            descriptive_outliers = clean.sub(med).abs().gt(max(6*1.4826*mad, 5))
            inventory.append(dict(original_label=label, model_input=label, sheet=sheet,
                excel_column=get_column_letter(position), frequency=frequency, group=group,
                raw_unit=unit, raw_definition=definition, transformation=rule,
                clean_unit='annual percentage rate' if rule == 'LEVEL' else ('percentage points of reported YoY growth' if rule == 'INDEX_MINUS_100' else '100 log points'),
                first_observation=str(valid.index.min()), last_observation=str(valid.index.max()),
                first_usable=str(usable.index.min()), last_usable=str(usable.index.max()),
                raw_observations=len(valid), usable_observations=len(usable),
                missing_raw=int(raw.isna().sum()), missing_transformed=int(clean.isna().sum()),
                zero_raw=int(raw.eq(0).sum()), negative_increment=int(intermediate.lt(0).sum()) if rule.startswith('DECUM') else 0,
                outlier_count=int(descriptive_outliers.sum()), annual_mean_shift_over_50pct=int(annual.gt(.5).sum()),
                publication_lag_days=lag, publication_lag_basis='ASSUMED reference period end + calendar days; not observed release',
                status='INCLUDED_SHORT_HISTORY' if len(usable)<36 else 'INCLUDED',
                late_start=valid.index.min() >= pd.Period('2024-01' if frequency=='M' else '2024Q1',freq=frequency),
                structural_missingness=bool(raw.iloc[:raw.first_valid_index().ordinal-raw.index[0].ordinal].isna().any()) if raw.first_valid_index() is not None else True,
                source_release_date=None, source_url=None,
                definition_status='LABEL_AND_PATTERN_INFERENCE; SOURCE_METADATA_UNVERIFIED'))
            for i, period in enumerate(raw.index):
                flags = []
                if pd.isna(raw.iloc[i]): flags.append('SOURCE_MISSING')
                if rule == 'INDEX_MOM_LOG' and raw.iloc[i] <= 0: flags.append('INVALID_NONPOSITIVE_INDEX')
                if rule.startswith('DECUM') and intermediate.iloc[i] <= 0: flags.append('NONPOSITIVE_FLOW')
                if pd.notna(raw.iloc[i]) and pd.isna(clean.iloc[i]): flags.append('UNUSABLE_TRANSFORM_OR_LAG_HISTORY')
                if descriptive_outliers.iloc[i]: flags.append('DESCRIPTIVE_OUTLIER_RETAINED')
                changes.append(dict(original_label=label, sheet=sheet, excel_cell=f'{get_column_letter(position)}{i+2}',
                    reference_period=str(period), frequency=frequency, raw_value=raw.iloc[i],
                    flow_or_input_value=intermediate.iloc[i], clean_value=clean.iloc[i],
                    transformation=rule, quality_flag=';'.join(flags)))
        transformed.append(clean_frame)
    concordance = None
    if master is not None:
        official = pd.read_parquet(master).set_index('quarter').gdp_real_yoy_pct
        actual = gdp.set_axis(gdp.index.astype(str))
        common = actual.index.intersection(official.index)
        discrepancy = float((actual.loc[common]-official.loc[common]).abs().max())
        concordance = dict(n_common=len(common), max_absolute_difference=discrepancy,
                           validated=bool(len(common)>0 and discrepancy < 1e-9))
        if not concordance['validated']:
            raise ValueError('Workbook GDP differs from official cumulative real YTD master')
    audit = dict(monthly_rows=len(m), quarterly_rows=len(q), monthly_predictors=m.shape[1],
        quarterly_predictors=predictors.shape[1], total_predictors=m.shape[1]+predictors.shape[1],
        gdp_label=target, gdp_convention='Published cumulative YTD real GDP YoY percent; index minus 100',
        gdp_master_concordance=concordance, duplicate_periods=0, duplicate_labels=0,
        workbook_source_metadata='Absent: no historical releases, source URLs, exact aggregation conventions or retrieval timestamps',
        information_class=INFORMATION_CLASS)
    return Inputs(m, q, transformed[0], transformed[1], gdp, pd.DataFrame(inventory), pd.DataFrame(changes), audit)


def release_gate(frame: pd.DataFrame, inventory: pd.DataFrame, origin, extra_days=0) -> pd.DataFrame:
    result = frame.copy()
    lags = inventory.set_index('original_label').publication_lag_days
    ends = frame.index.to_timestamp(how='end').normalize()
    for c in result:
        result.loc[ends+pd.Timedelta(days=int(lags[c])+extra_days)>pd.Timestamp(origin), c] = np.nan
    return result

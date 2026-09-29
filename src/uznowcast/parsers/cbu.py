"""CBU activation-date observations, without calendar-day forward filling."""
from io import BytesIO
import re
import openpyxl
import numpy as np
import pandas as pd


def parse_fx(payload, currency='USD') -> pd.DataFrame:
    if not isinstance(payload, list) or not payload:
        raise ValueError('CBU expected non-empty JSON list')
    records = []
    for item in payload:
        if not isinstance(item, dict) or not {'Ccy', 'Date', 'Rate', 'Nominal'} <= item.keys():
            raise ValueError('CBU schema missing required fields')
        if item['Ccy'] != currency:
            continue
        date = pd.to_datetime(item['Date'], format='%d.%m.%Y', errors='raise')
        rate, nominal = float(item['Rate']), float(item['Nominal'])
        if not np.isfinite([rate, nominal]).all() or rate <= 0 or nominal <= 0:
            raise ValueError('CBU nonpositive/nonfinite rate or nominal')
        records.append(dict(reference_date=date, reference_period=date.strftime('%Y-%m-%d'),
                            frequency='D', raw_value=rate, nominal=nominal,
                            normalized_daily=rate / nominal, clean_value=rate / nominal,
                            quality_flag='', clean_unit=f'UZS per {currency}'))
    if not records:
        raise ValueError(f'CBU missing exact currency: {currency}')
    return deduplicate_daily(pd.DataFrame(records))


def deduplicate_daily(frame: pd.DataFrame) -> pd.DataFrame:
    for _, group in frame.groupby('reference_date'):
        if len(group[['raw_value', 'nominal']].drop_duplicates()) != 1:
            raise ValueError('Conflicting duplicate CBU activation date')
    return frame.sort_values('reference_date', kind='stable').drop_duplicates('reference_date').reset_index(drop=True)


def parse_m2_xlsx(content: bytes, row: dict, contract: dict) -> pd.DataFrame:
    """Extract the exact broad-money row from the CBU DCS workbook."""
    if row['row_field_selector'] != contract['registry_selector']:
        raise ValueError('Registry M2 selector changed; review CBU adapter')
    if row['raw_unit'] != contract['registry_unit'] or row['rule_codes'] != contract['rule_codes']:
        raise ValueError('Registry M2 unit/transformation changed; review CBU adapter')
    workbook = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        if workbook.sheetnames != [contract['sheet']]:
            raise ValueError(f'CBU M2 workbook sheets changed: {workbook.sheetnames}')
        sheet = workbook[contract['sheet']]
        if sheet['A1'].value != contract['title'] or sheet['A2'].value != contract['unit']:
            raise ValueError('CBU M2 workbook title/unit changed')
        labels = [(cells[0].row, cells[0].value)
                  for cells in sheet.iter_rows(min_col=1, max_col=1)
                  if isinstance(cells[0].value, str)]
        matches = [number for number, label in labels if label.strip() == contract['label']]
        if len(matches) != 1:
            raise ValueError(f'CBU M2 exact row selector matched {len(matches)} rows')
        selected_row = matches[0]
        records, seen = [], set()
        for column in range(2, sheet.max_column + 1):
            label = sheet.cell(contract['header_row'], column).value
            value = sheet.cell(selected_row, column).value
            if label is None:
                if value is not None:
                    raise ValueError('CBU M2 value has no period header')
                continue
            if not isinstance(label, str) or not re.fullmatch(r'\d{4}-(?:0[1-9]|1[0-2])', label):
                raise ValueError(f'CBU M2 unexpected period header: {label!r}')
            if label in seen:
                raise ValueError(f'CBU M2 duplicate period: {label}')
            seen.add(label)
            if not isinstance(value, (int, float)) or isinstance(value, bool) or not np.isfinite(value):
                raise ValueError(f'CBU M2 nonnumeric value: {label}: {value!r}')
            if value <= 0:
                raise ValueError(f'CBU M2 nonpositive stock: {label}: {value!r}')
            period = pd.Period(label, freq='M')
            records.append(dict(reference_period=str(period),
                                reference_date=period.to_timestamp(how='end').normalize(),
                                frequency='M', raw_value=float(value), clean_value=np.nan,
                                quality_flag='', selected_label=contract['label']))
        if not records:
            raise ValueError('CBU M2 selected row has no observations')
        periods = pd.PeriodIndex([record['reference_period'] for record in records], freq='M')
        expected = pd.period_range(periods.min(), periods.max(), freq='M')
        if not periods.equals(expected):
            raise ValueError('CBU M2 monthly headers are not contiguous and ordered')
        return pd.DataFrame(records)
    finally:
        workbook.close()

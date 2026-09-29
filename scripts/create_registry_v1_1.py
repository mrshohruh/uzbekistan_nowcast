"""Create the approved V1.1 registry while retaining V1.0 unchanged."""
from __future__ import annotations

from copy import copy
from pathlib import Path
import hashlib
import json
import shutil

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'registry/uzbekistan_nowcasting_v1_registry.xlsx'
ARCHIVE = ROOT / 'registry/archive/uzbekistan_nowcasting_v1_registry_v1.0.xlsx'
TARGET = ROOT / 'registry/uzbekistan_nowcasting_v1.1_registry.xlsx'
EVIDENCE = ROOT / 'registry/uzbekistan_nowcasting_v1.1_changes.json'

APPROVED = {
    ('construction', 'Verified end'): ('2026-M07', '2026-M08'),
    ('construction', 'Lag basis / release convention'):
        ('Jul-2026 source updated 2026-08-27',
         'Aug-2026 source updated 2026-09-28'),
    ('cpi_headline', 'Row / field selector'):
        ('Composite index', 'Code=1; Klassifikator_en=Cumulative index'),
    ('m2', 'Machine/download URL'):
        ('https://cbu.uz/upload/open_data/0020/4-009-0020_eng.json',
         'https://cbu.uz/sdmx/public/DCS_Uzbekistan_Online.xlsx'),
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def workbook_values(path: Path) -> dict[tuple[str, str], object]:
    workbook = openpyxl.load_workbook(path, read_only=False, data_only=False)
    try:
        return {(sheet.title, cell.coordinate): cell.value
                for sheet in workbook.worksheets
                for row in sheet.iter_rows()
                for cell in row}
    finally:
        workbook.close()


def main() -> None:
    ARCHIVE.parent.mkdir(parents=True, exist_ok=True)
    if ARCHIVE.exists() and ARCHIVE.read_bytes() != SOURCE.read_bytes():
        raise RuntimeError('Existing archived V1.0 differs from the authoritative source')
    if not ARCHIVE.exists():
        shutil.copyfile(SOURCE, ARCHIVE)

    workbook = openpyxl.load_workbook(SOURCE)
    registry = workbook['V1 Registry']
    headers = {cell.value: cell.column for cell in registry[4]}
    rows = {registry.cell(row, headers['Variable key']).value: row
            for row in range(5, registry.max_row + 1)}
    changes = []
    for (key, field), (expected, replacement) in APPROVED.items():
        cell = registry.cell(rows[key], headers[field])
        if cell.value != expected:
            raise RuntimeError(f'Unexpected V1.0 value for {key}/{field}: {cell.value!r}')
        changes.append(dict(variable_key=key, field=field, old_value=expected,
                            new_value=replacement))
        cell.value = replacement
    version = workbook['README']['B3']
    if version.value != 'V1.0 — verified 2026-09-29':
        raise RuntimeError(f'Unexpected V1.0 version label: {version.value!r}')
    version.value = 'V1.1 — verified 2026-09-29'
    workbook.save(TARGET)
    workbook.close()

    before, after = workbook_values(SOURCE), workbook_values(TARGET)
    expected_changed_cells = {('README', 'B3')}
    for key, field in APPROVED:
        expected_changed_cells.add(('V1 Registry',
                                    openpyxl.utils.get_column_letter(headers[field]) + str(rows[key])))
    actual_changed_cells = {cell for cell in before if before[cell] != after[cell]}
    if actual_changed_cells != expected_changed_cells:
        raise RuntimeError(f'Unapproved registry cell changes: {actual_changed_cells ^ expected_changed_cells}')

    payload = dict(
        old_registry=SOURCE.relative_to(ROOT).as_posix(),
        archived_v1_0=ARCHIVE.relative_to(ROOT).as_posix(),
        active_v1_1=TARGET.relative_to(ROOT).as_posix(),
        v1_0_sha256=sha256(SOURCE),
        archived_v1_0_sha256=sha256(ARCHIVE),
        v1_1_sha256=sha256(TARGET),
        verification_date='2026-09-29',
        approved_changes=changes,
        changed_cells=sorted(f'{sheet}!{cell}' for sheet, cell in actual_changed_cells),
    )
    EVIDENCE.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')


if __name__ == '__main__':
    main()

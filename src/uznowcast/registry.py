"""Workbook contract loading. Original labels are retained alongside normalized keys."""
from dataclasses import dataclass
from pathlib import Path
import json
import re
from urllib.parse import urlparse

import openpyxl

PHASE2A = ("gdp_real_yoy", "industrial_production", "usd_uzs")
PILOT8 = ("gdp_real_yoy", "industrial_production", "construction", "cpi_headline",
          "exports_total", "imports_total", "usd_uzs", "m2")
# Backward-compatible name used by Phase 2A callers and tests.
PILOT = PHASE2A
STATUSES = {"READY_API", "READY_PAGE", "ARCHIVE_SPIDER", "DERIVED", "PROXY", "EXTERNAL"}
FREQUENCIES = {"Quarterly", "Monthly", "Daily", "Monthly archive / current JSON snapshot"}
HEADERS = (
    "V1 order", "Variable key", "Block", "Display name", "Provider",
    "Native indicator / dataset ID", "Dataset/page ID", "Row / field selector",
    "Human source URL", "Machine/download URL", "Native frequency", "Verified start",
    "Verified end", "Raw unit", "Flow / stock type", "Cumulative flag",
    "Typical publication lag (days)", "Lag basis / release convention",
    "Required transformation", "Clean model field", "Revision characteristics",
    "Structural breaks / caveats", "Model role", "Automation status",
)


def normalize(label: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", label.lower()).strip("_")


@dataclass
class Registry:
    rows: dict[str, dict]
    original_labels: dict[str, str]
    transform_rules: dict[str, dict]
    path: Path
    version: str
    verification_date: str

    def pilot(self) -> list[dict]:
        return self.scope('pilot')

    def scope(self, name: str) -> list[dict]:
        keys = {'pilot': PHASE2A, 'pilot8': PILOT8, 'v1': tuple(self.rows)}.get(name)
        if keys is None:
            raise ValueError(f'Unknown build scope: {name}')
        return sorted((self.rows[k] for k in keys), key=lambda r: r['v1_order'])


def load_registry(path: Path, contracts_path: Path | None = None) -> Registry:
    path = Path(path)
    contracts_path = contracts_path or path.parent.parent / 'config/registry_contracts.json'
    contracts = json.loads(contracts_path.read_text(encoding='utf-8'))
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook['V1 Registry']
        headers = [c.value for c in sheet[4]]
        if headers != list(HEADERS):
            raise ValueError(f'Registry row-4 header mismatch: {headers}')
        labels = {normalize(h): h for h in headers}
        rule_sheet = workbook['Transform Rules']
        rule_headers = [normalize(c.value) for c in rule_sheet[3]]
        rules = {r[0]: dict(zip(rule_headers, r)) for r in rule_sheet.iter_rows(min_row=4, values_only=True) if r[0]}
        if set(rules) != set(contracts['rule_codes']):
            raise ValueError('Unrecognized Transform Rules codes')
        for code, rule in rules.items():
            if rule['exact_rule'] != contracts['rule_text'][code]:
                raise ValueError(f'Transform rule changed; review implementation: {code}')
        rows = {}
        clean_fields = set()
        orders = set()
        for values in sheet.iter_rows(min_row=5, values_only=True):
            if not any(v is not None for v in values):
                continue
            row = dict(zip(labels, values))
            required = ['variable_key', 'v1_order', 'provider', 'native_frequency', 'raw_unit',
                        'required_transformation', 'clean_model_field', 'automation_status',
                        'row_field_selector', 'flow_stock_type', 'cumulative_flag', 'model_role']
            if any(row[k] is None or str(row[k]).strip() == '' for k in required):
                raise ValueError(f'Missing mandatory registry fields: {row.get("variable_key")}')
            key = row['variable_key']
            if key in rows or not re.fullmatch(r'[a-z][a-z0-9_]*', key):
                raise ValueError(f'Duplicate/invalid variable key: {key}')
            if row['automation_status'] not in STATUSES or row['native_frequency'] not in FREQUENCIES:
                raise ValueError(f'Unrecognized status/frequency: {key}')
            if row['clean_model_field'] in clean_fields or row['v1_order'] in orders:
                raise ValueError(f'Duplicate clean field/order: {key}')
            if not re.fullmatch(r'[a-z][a-z0-9_]*', row['clean_model_field']):
                raise ValueError(f'Invalid clean model field: {key}')
            if not isinstance(row['v1_order'], int) or row['v1_order'] < 1:
                raise ValueError(f'Invalid registry order: {key}')
            clean_fields.add(row['clean_model_field'])
            orders.add(row['v1_order'])
            text = row['required_transformation']
            if text not in contracts['transformations']:
                raise ValueError(f'Unrecognized transformation: {key}: {text}')
            row['rule_codes'] = contracts['transformations'][text]
            if row['automation_status'] != 'DERIVED':
                if not row['native_indicator_dataset_id'] or not row['dataset_page_id']:
                    raise ValueError(f'Missing source identifier: {key}')
                urls = [row['human_source_url'], row['machine_download_url']]
                if not any(isinstance(u, str) and urlparse(u).scheme == 'https' for u in urls):
                    raise ValueError(f'Missing official source URL: {key}')
                if row['automation_status'] == 'READY_API' and not str(row['machine_download_url']).startswith('https://'):
                    raise ValueError(f'Missing API URL: {key}')
            rows[key] = row
        if not set(PILOT8) <= rows.keys():
            raise ValueError('Phase 2B pilot keys missing')
        version_label = workbook['README']['B3'].value
        version_match = re.fullmatch(r'(V\d+\.\d+) — verified (\d{4}-\d{2}-\d{2})', str(version_label))
        if not version_match:
            raise ValueError(f'Invalid registry version label: {version_label!r}')
        return Registry(rows, labels, rules, path, version_match.group(1), version_match.group(2))
    finally:
        workbook.close()

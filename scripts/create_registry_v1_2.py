"""Create the approved V1.2 registry while retaining V1.0 and V1.1 unchanged.

Applies the source-maintenance corrections documented in
``docs/phase2c_registry_recommendations.md``. The ``russia_ipi`` recommendation
is deliberately excluded per Phase 3B instructions: provider stays Rosstat and
no TLS policy is changed here. The retail/wholesale row-selector recommendations
were verified to already match the live SIAT label in V1.1 and therefore require
no cell change; that verification is recorded in the changelog.
"""
from __future__ import annotations

from pathlib import Path
import hashlib
import json
import shutil

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
V10 = ROOT / 'registry/archive/uzbekistan_nowcasting_v1_registry_v1.0.xlsx'
SOURCE = ROOT / 'registry/uzbekistan_nowcasting_v1.1_registry.xlsx'
ARCHIVE_V11 = ROOT / 'registry/archive/uzbekistan_nowcasting_v1.1_registry.xlsx'
TARGET = ROOT / 'registry/uzbekistan_nowcasting_v1.2_registry.xlsx'
EVIDENCE = ROOT / 'registry/uzbekistan_nowcasting_v1.2_changes.json'

PPI_NEW_BREAKS = (
    "since April 2024 the scope of reporting entities has been expanded, "
    "price indices for January, February, March 2024 have been recalculated. "
    "Period labels mix Cyrillic 'М' (U+041C) for 2016-M01..2020-M12 with Latin "
    "'M' elsewhere; parser normalizes to Latin without rewriting the raw payload."
)

APPROVED = {
    ('manufacturing', 'Row / field selector'): (
        'Manufacturing',
        'Code=C; Klassifikator_en=Manufacturing industry'),
    ('electricity_gas', 'Row / field selector'): (
        'Electricity, gas, steam and air conditioning supply',
        'Code=D; Klassifikator_en=Electricity, gas, steam and air conditioning'),
    ('retail_trade', 'Raw unit'): (
        'billion UZS',
        'billion UZS (source: million sums; standardization scale 0.001)'),
    ('wholesale_trade', 'Raw unit'): (
        'billion UZS',
        'billion UZS (source: million sums; standardization scale 0.001)'),
    ('gold_exports_proxy', 'Row / field selector'): (
        'Other goods (gold-dominated residual category)',
        'Code=9; Klassifikator_en=Other goods (gold-dominated residual category)'),
    ('fx_reserves_ex_gold', 'Machine/download URL'): (
        'https://cbu.uz/upload/open_data/0017/4-009-0017_eng.json',
        'https://cbu.uz/en/statistics/e-gdds/data/111574/'),
    ('interbank_payments', 'Raw unit'): (
        'UZS (transaction amount)',
        'billion UZS (source: thousand UZS; standardization scale 1e-6)'),
    ('ppi', 'Structural breaks / caveats'): (
        'Monitor producer-price methodology/rebasing notices.',
        PPI_NEW_BREAKS),
}

# Verified in this environment: V1.1 already carries the live SIAT selectors
# 'Republic of Uzbekistan / retail trade turnover' and
# 'Republic of Uzbekistan / wholesale trade turnover'. The Phase 2C
# recommendation to adopt the exact SIAT label is therefore satisfied and no
# cell edit is required.
VERIFIED_NO_OP = {
    'retail_trade': ('Row / field selector', 'Republic of Uzbekistan / retail trade turnover'),
    'wholesale_trade': ('Row / field selector', 'Republic of Uzbekistan / wholesale trade turnover'),
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def workbook_values(path: Path) -> dict[tuple[str, str], object]:
    workbook = openpyxl.load_workbook(path, read_only=False, data_only=False)
    try:
        return {(sheet.title, cell.coordinate): cell.value
                for sheet in workbook.worksheets
                for row in sheet.iter_rows()
                for cell in row}
    finally:
        workbook.close()


def main() -> None:
    if not V10.exists() or not SOURCE.exists():
        raise RuntimeError('V1.0 archive and/or V1.1 workbook missing; run create_registry_v1_1.py first')
    ARCHIVE_V11.parent.mkdir(parents=True, exist_ok=True)
    if ARCHIVE_V11.exists() and ARCHIVE_V11.read_bytes() != SOURCE.read_bytes():
        raise RuntimeError('Existing archived V1.1 differs from the authoritative V1.1 workbook')
    if not ARCHIVE_V11.exists():
        shutil.copyfile(SOURCE, ARCHIVE_V11)

    workbook = openpyxl.load_workbook(SOURCE)
    registry = workbook['V1 Registry']
    headers = {cell.value: cell.column for cell in registry[4]}
    rows = {registry.cell(row, headers['Variable key']).value: row
            for row in range(5, registry.max_row + 1)}
    changes = []
    for (key, field), (expected, replacement) in APPROVED.items():
        cell = registry.cell(rows[key], headers[field])
        if cell.value != expected:
            raise RuntimeError(f'Unexpected V1.1 value for {key}/{field}: {cell.value!r}')
        changes.append(dict(variable_key=key, field=field, old_value=expected,
                            new_value=replacement))
        cell.value = replacement
    for key, (field, expected) in VERIFIED_NO_OP.items():
        cell = registry.cell(rows[key], headers[field])
        if cell.value != expected:
            raise RuntimeError(f'Unexpected V1.1 value for verified no-op {key}/{field}: {cell.value!r}')
    version = workbook['README']['B3']
    if version.value != 'V1.1 — verified 2026-09-29':
        raise RuntimeError(f'Unexpected V1.1 version label: {version.value!r}')
    version.value = 'V1.2 — verified 2026-09-29'
    workbook.save(TARGET)
    workbook.close()

    before, after = workbook_values(SOURCE), workbook_values(TARGET)
    expected_changed_cells = {('README', 'B3')}
    for key, field in APPROVED:
        expected_changed_cells.add(('V1 Registry',
                                    openpyxl.utils.get_column_letter(headers[field]) + str(rows[key])))
    actual_changed_cells = {cell for cell in before if before[cell] != after[cell]}
    if actual_changed_cells != expected_changed_cells:
        raise RuntimeError(f'Unapproved registry cell changes: {actual_changed_cells ^ expected_changed_cells}')

    payload = dict(
        base_registry=SOURCE.relative_to(ROOT).as_posix(),
        archived_v1_0=V10.relative_to(ROOT).as_posix(),
        archived_v1_1=ARCHIVE_V11.relative_to(ROOT).as_posix(),
        active_v1_2=TARGET.relative_to(ROOT).as_posix(),
        v1_0_sha256=sha256(V10),
        v1_1_sha256=sha256(SOURCE),
        archived_v1_1_sha256=sha256(ARCHIVE_V11),
        v1_2_sha256=sha256(TARGET),
        verification_date='2026-09-29',
        approved_changes=changes,
        verified_no_op=[dict(variable_key=key, field=field, value=value)
                        for key, (field, value) in VERIFIED_NO_OP.items()],
        changed_cells=sorted(f'{sheet}!{cell}' for sheet, cell in actual_changed_cells),
    )
    EVIDENCE.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')


if __name__ == '__main__':
    main()

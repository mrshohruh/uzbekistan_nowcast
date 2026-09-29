"""Strict parsers for external registry-approved providers."""
from io import BytesIO
import re

import numpy as np
import openpyxl
import pandas as pd

from uznowcast.io.external import world_bank_release_timestamp


def parse_world_bank_gold(content: bytes, row: dict) -> tuple[pd.DataFrame, str | None]:
    if row['row_field_selector'] != 'Commodity=Gold' or row['raw_unit'] != 'USD per troy ounce':
        raise ValueError('World Bank gold registry contract changed')
    workbook = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        if 'Monthly Prices' not in workbook.sheetnames:
            raise ValueError('Pink Sheet Monthly Prices worksheet missing')
        sheet = workbook['Monthly Prices']
        if sheet['A1'].value != 'World Bank Commodity Price Data (The Pink Sheet)':
            raise ValueError('Pink Sheet title changed')
        release = world_bank_release_timestamp(str(sheet['A4'].value))
        headers = [cell.value for cell in sheet[5]]
        matches = [index + 1 for index, value in enumerate(headers) if value == 'Gold']
        if len(matches) != 1 or sheet.cell(6, matches[0]).value != '($/troy oz)':
            raise ValueError('Pink Sheet exact Gold column/unit changed')
        column = matches[0]
        records = []
        for cells in sheet.iter_rows(min_row=7, max_col=column):
            label = cells[0].value
            if label is None:
                break
            if not isinstance(label, str) or not re.fullmatch(r'\d{4}M(?:0[1-9]|1[0-2])', label):
                raise ValueError(f'Pink Sheet unexpected monthly label: {label!r}')
            value = cells[column - 1].value
            if not isinstance(value, (int, float)) or isinstance(value, bool) or not np.isfinite(value) or value <= 0:
                raise ValueError(f'Pink Sheet invalid Gold value: {label}: {value!r}')
            period = pd.Period(label.replace('M', '-'), freq='M')
            records.append(dict(reference_period=str(period),
                                reference_date=period.to_timestamp(how='end').normalize(),
                                frequency='M', raw_value=float(value), clean_value=np.nan,
                                quality_flag='', clean_unit='percent log change'))
        return pd.DataFrame(records), release
    finally:
        workbook.close()
